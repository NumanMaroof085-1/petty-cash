"""Preserve existing Word structure while correcting confirmed schema facts.

Writes candidates to documentation/corrected; originals remain unchanged until
rendering is verified. Run with the selected workspace Python runtime.
"""
from pathlib import Path
from copy import deepcopy
import importlib.util
import shutil
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "documentation" / "corrected"
BACKUP = ROOT / "documentation" / "backups" / "2026-10-12"
HELPER = Path("C:/Users/numan/.codex/plugins/cache/openai-primary-runtime/documents/26.1011.11113/skills/documents/scripts/comments_add.py")


def replace(paragraph, old, new):
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    return True


def paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def append_notes(doc, heading, notes):
    doc.add_heading(heading, level=1)
    for note in notes:
        doc.add_paragraph(note)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=True)
    spec = importlib.util.spec_from_file_location("document_comments", HELPER)
    helper = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helper)
    names = ["Petty Cash App - Database Schema.docx", "Petty Cash Request System - Project Overview.docx", "Configuration document.docx"]
    for name in names:
        source = ROOT / name
        backup = BACKUP / name
        if not backup.exists():
            shutil.copy2(source, backup)
        doc = Document(source)
        comments = []
        if name.startswith("Petty Cash App"):
            changed = 0
            for p in paragraphs(doc):
                if replace(p, "Currency (Calculated)", "Currency (Simple)"):
                    changed += 1
                    comments.append(("Currency (Simple)", "Corrected from Calculated: the live Total Price column was confirmed Simple. The client writes and verifies Quantity multiplied by Unit Price."))
            assert changed == 1, "Expected exactly one Total Price type correction"
            # Add the confirmed existing Other Category row to the Line Item table.
            table = next(t for t in doc.tables if any("hr_totalprice" in c.text for r in t.rows for c in r.cells))
            if not any("hr_othercategory" in c.text for r in table.rows for c in r.cells):
                # Existing document grids contain fractional twip widths that
                # python-docx cannot use in add_row; clone a formatted row instead.
                table._tbl.append(deepcopy(table.rows[-1]._tr))
                cells = table.rows[-1].cells
                for cell, text in zip(cells, ["Other Category", "hr_othercategory", "Single Line of Text", "Details when Category is Others. Stored on each Line Item, not the Request header."]):
                    cell.text = text
            append_notes(doc, "Confirmed balance rules and proposed schema addition", [
                "The existing Cash Transaction table is the balance ledger. Its logical name is hr_cashtransaction and its entity set is hr_cashtransactions. No additional balance table is required.",
                "Transaction types: Advance Given 123140000 adds to balance; Expense Approved 123140001 subtracts; Cash Returned 123140002 subtracts. Amounts are positive. Reimbursement uses Advance Given. Everyone starts at zero.",
                "Positive balance means company cash held by the employee. Negative means the company owes the employee. Both approved Own Pocket and Company Cash requests subtract once. Pending and rejected requests do not affect balance.",
                "Contact hr_currentbalance is a server-maintained summary; Cash Transactions remain the source of truth. The new website pages sum the employee ledger and do not directly overwrite the summary.",
                "Proposed, not yet created: hr_postingkey (schema hr_PostingKey), Text, maximum 100 characters, with an active unique alternate key. Accounts operations use a UUID; approved expenses use expense-<request-guid>. The existing autonumber reference does not prevent duplicate expense postings.",
                "Receipt upload is mandatory in the client. hr_receiptimage stores one file per row; multiple attachments require related attachment records and are not implemented by this version.",
                "Cash Returned means physical cash returned by the employee, not arbitrary manual balance correction. Approved financial history must not be silently edited or deleted."
            ])
        elif name.startswith("Petty Cash Request"):
            replacements = {
                "Automated Processing: Power Automate triggers immediately in the background, identifies the submitting employee, and links their department and assigned manager.": "Automated Processing: approval starts only after the complete request, all intended Line Items, and mandatory receipt have been validated on the server. The portal already reads the employee profile and saves Department and Assigned Manager lookups.",
                "4. Planned Module: Cash Advance Balance Tracking": "4 Signed Cash Balance Tracking",
                "Purpose: To provide logged-in employees with a real-time view of their available company cash advance balance directly on the request form.": "Purpose: To show the employee's signed cash balance on the request form. Positive means company cash held; negative means the company owes the employee. It is not a spending limit.",
                "Liquid or front-end JavaScript Web API will fetch the user's available balance from their Dataverse contact record or dedicated balance ledger table (hr_cashadvancebalance).": "The shared JavaScript reads and sums the employee's existing Cash Transaction ledger (hr_cashtransaction / hr_cashtransactions). Contact hr_currentbalance remains a server-maintained summary. No hr_cashadvancebalance table is required.",
                "Once a request against a cash advance is approved, Power Automate will automatically update and deduct the approved amount from the user's balance in Dataverse.": "Once either a Company Cash or Own Pocket request is approved, the trusted backend records one Expense Approved transaction linked to that request. Cash given, including reimbursement, increases balance; cash returned decreases it. Everyone starts at zero.",
                "A new Power Pages Web Template and List view (\"My Petty Cash Requests\") will be built.": "My Petty Cash Requests.html provides the new employee dashboard through session-authenticated Web API reads, with balance, saved request status, and manager comments. Its website page/template and shared JavaScript Web File still need deployment.",
                "Live status updates triggered by Power Automate approval outcomes will reflect dynamically on this dashboard.": "Refresh retrieves the latest saved approval outcomes and comments. Power Pages caching can delay changes made by another process; this is not push-based real-time streaming."
            }
            for p in paragraphs(doc):
                for old, new in replacements.items():
                    replace(p, old, new)
            comments.append(("The shared JavaScript reads and sums", "Updated to use the existing Cash Transaction table and the confirmed signed-balance rules. No extra balance table is required."))
            append_notes(doc, "Accounts cash management and implementation status", [
                "Accounts Cash Management.html is a new local page fragment for recording cash given to an employee or physically returned by an employee. The employee is selected by Contact GUID; entered amounts are movements, not replacement balances.",
                "The page and posting flow must be restricted to the Petty Cash Accounts web role. An hr_userrole Choice or a hidden navigation item does not enforce authorization. Employees may read only their own requests and cash transactions.",
                "Local implementation: the request-form balance widget, employee dashboard, accounts form, shared pettycash-modules.js, and regression tests. Cash posting is disabled by default until its authenticated backend is configured and tested.",
                "Tenant work pending: page/template and Web File deployment, explicit API field lists, scoped permissions and Accounts web role, proposed hr_postingkey column/key, secured posting flow, server-side completeness/approval processing, and concurrency-safe summary maintenance. No live configuration or financial posting has been performed by the local code changes.",
                "Business Central integration follows a verified approval and cash-posting workflow. IMPLEMENTATION_LOG.md records every new artifact and configuration requirement; BALANCE_MODULE_SETUP.md contains deployment and backend contracts."
            ])
        else:
            for p in paragraphs(doc):
                replace(p, "Department (hr_department)", "Department (hr_departments)")
                replace(p, "Webapi/hr_department/", "Webapi/hr_departments/")
                replace(p, "* (Must include hr_receiptimage)", "Explicit field list in WEBAPI_SITE_SETTINGS.md, including hr_receiptimage")
                if p.text.strip() == "*":
                    replace(p, "*", "Explicit field list in WEBAPI_SITE_SETTINGS.md")
                replace(p, "Exposes columns and explicitly unblocks image file uploads.", "Exposes explicitly listed Request columns, including the receipt upload field.")
            # Preserve the recorded historical settings but label them as historical.
            for p in doc.paragraphs:
                replace(p, "2. Dataverse Table Permissions", "2 Historical Table Permissions")
                replace(p, "To allow logged-in users to submit forms and the system to establish background lookups without throwing 403 Forbidden or 500 Internal Server errors, the following permissions were mapped to the Authenticated Users web role.", "The table below records the original troubleshooting permissions. Global employee permissions are not the production security model. Apply the scoped permissions below before publishing employee balances and requests.")
            comments.append(("Historical Table Permissions", "The original Global permissions are preserved as history, not recommended production access. See the new scoped-permission section."))
            append_notes(doc, "Balance module website configuration", [
                "Publish pettycash-modules.js as the root Web File /pettycash-modules.js. Deploy the updated request page, My Petty Cash Requests.html, and Accounts Cash Management.html. Create the corresponding pages/templates and navigation entries in Power Pages; local HTML does not create tenant records.",
                "Add explicit Contact fields contactid, fullname, hr_currentbalance, statecode while preserving _hr_department_value and _hr_reportstomanager_value. Add hr_managercomments to the existing Request allow-list. Add Webapi/hr_cashtransaction/enabled=true and its explicit field list from WEBAPI_SITE_SETTINGS.md.",
                "Employee permissions: Self Contact Read, Contact-scoped Request Read through Submitted By, Contact-scoped Cash Transaction Read through Employee, and related Line Item access through Request. Verify required submission privileges and remove overlapping Global employee reads. Employees cannot write balances, approval outcomes, or ledger rows.",
                "Accounts permissions: a dedicated Petty Cash Accounts web role, restricted page access, authorized Contact and Cash Transaction reads, and access to only the secured accounts posting flow. No portal direct balance or ledger Write/Create/Delete is needed by the accounts page.",
                "Use a solution-aware When Power Pages calls a flow trigger attached to the site with only the Accounts role. Configure the generated flow GUID in PC_MODULE_CONFIG.accountsFlowId; enable accountsSecurityReviewed only after backend validation and duplicate/concurrency tests. Never expose a public signed HTTP trigger URL.",
                "Create the proposed Cash Transaction hr_postingkey Text column and activate its unique alternate key before posting. Keep amounts positive and derive the sign from the confirmed transaction type. Server writes must be duplicate-safe and concurrency-safe; a simple read-old-balance then Update Contact is insufficient.",
                "The accounts flow, approval flow, new schema key, roles, and website settings are requirements, not tenant changes performed by this local implementation. Detailed configuration and verification are in BALANCE_MODULE_SETUP.md; every subsequent change must also update IMPLEMENTATION_LOG.md."
            ])
        raw = OUT / ("raw-" + name)
        doc.save(raw)
        final = OUT / name
        helper.add_comments(str(raw), str(final), comments, author="Project documentation", ignore_case=False, require_all=True)
        print("Prepared", final)


if __name__ == "__main__":
    main()
